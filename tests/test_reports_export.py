"""Tests for report export functionality."""

import pytest
from decimal import Decimal
from datetime import datetime, date
from django.utils import timezone
from openpyxl import load_workbook
from io import BytesIO

from src.reports.export import (
    export_account_balance_report,
    export_transaction_history_report,
    export_trial_balance_report,
    export_balance_sheet_report,
    export_loan_aging_report,
    get_excel_response,
)


class TestAccountBalanceExport:
    """Test account balance report export."""
    
    def test_export_account_balance_report(self):
        """Test generating account balance Excel report."""
        account_data = {
            'account_code': 'AC-001',
            'account_name': 'Cash',
            'account_type': 'ASSET',
            'balance': Decimal('5000.00'),
            'debit_total': Decimal('10000.00'),
            'credit_total': Decimal('5000.00'),
            'as_of_date': datetime.now(),
        }
        
        wb = export_account_balance_report(account_data)
        
        assert wb is not None
        assert len(wb.sheetnames) >= 1


class TestTransactionHistoryExport:
    """Test transaction history report export."""
    
    def test_export_transaction_history_report(self):
        """Test generating transaction history Excel report."""
        export_data = {
            'account_code': 'AC-001',
            'account_name': 'Cash',
            'account_type': 'ASSET',
            'transactions': [
                {
                    'transaction_id': 1,
                    'date': datetime.now(),
                    'idempotency_key': 'idem-001',
                    'description': 'Test transaction',
                    'entry_type': 'DEBIT',
                    'debit': Decimal('100.00'),
                    'credit': Decimal('0.00'),
                    'running_balance': Decimal('100.00'),
                    'status': 'POSTED',
                }
            ],
        }
        
        wb = export_transaction_history_report(export_data)
        
        assert wb is not None
        assert len(wb.sheetnames) >= 1


class TestTrialBalanceExport:
    """Test trial balance report export."""
    
    def test_export_trial_balance_report(self):
        """Test generating trial balance Excel report."""
        export_data = {
            'accounts': [
                {
                    'account_type': 'ASSET',
                    'account_code': 'AC-001',
                    'account_name': 'Cash',
                    'debit_balance': Decimal('5000.00'),
                    'credit_balance': Decimal('0.00'),
                },
                {
                    'account_type': 'LIABILITY',
                    'account_code': 'LI-001',
                    'account_name': 'Accounts Payable',
                    'debit_balance': Decimal('0.00'),
                    'credit_balance': Decimal('5000.00'),
                }
            ],
            'total_debits': Decimal('5000.00'),
            'total_credits': Decimal('5000.00'),
            'is_balanced': True,
            'as_of_date': datetime.now(),
        }
        
        wb = export_trial_balance_report(export_data)
        
        assert wb is not None
        assert len(wb.sheetnames) >= 1


class TestBalanceSheetExport:
    """Test balance sheet report export."""
    
    def test_export_balance_sheet_report(self):
        """Test generating balance sheet Excel report."""
        export_data = {
            'assets': [
                {
                    'account_code': 'AC-001',
                    'account_name': 'Cash',
                    'balance': Decimal('5000.00'),
                }
            ],
            'liabilities': [
                {
                    'account_code': 'LI-001',
                    'account_name': 'Accounts Payable',
                    'balance': Decimal('3000.00'),
                }
            ],
            'equity': [
                {
                    'account_code': 'EQ-001',
                    'account_name': 'Retained Earnings',
                    'balance': Decimal('2000.00'),
                }
            ],
            'total_assets': Decimal('5000.00'),
            'total_liabilities': Decimal('3000.00'),
            'total_equity': Decimal('2000.00'),
            'is_balanced': True,
            'as_of_date': datetime.now(),
        }
        
        wb = export_balance_sheet_report(export_data)
        
        assert wb is not None
        assert len(wb.sheetnames) >= 1


class TestLoanAgingExport:
    """Test loan aging report export."""
    
    def test_export_loan_aging_report(self):
        """Test generating loan aging Excel report."""
        export_data = {
            'summary': [
                {
                    'bucket_name': 'Current',
                    'days_range': '0-29 days',
                    'loan_count': 5,
                    'total_outstanding': Decimal('50000.00'),
                    'total_principal': Decimal('45000.00'),
                    'total_interest': Decimal('5000.00'),
                },
            ],
            'total_outstanding': Decimal('50000.00'),
            'loan_details': [],
            'as_of_date': datetime.now(),
        }
        
        wb = export_loan_aging_report(export_data)
        
        assert wb is not None
        assert len(wb.sheetnames) >= 1


class TestGetExcelResponse:
    """Test Excel response generation."""
    
    def test_get_excel_response(self):
        """Test generating HTTP response for Excel file."""
        account_data = {
            'account_code': 'AC-001',
            'account_name': 'Cash',
            'account_type': 'ASSET',
            'balance': Decimal('5000.00'),
            'debit_total': Decimal('10000.00'),
            'credit_total': Decimal('5000.00'),
            'as_of_date': datetime.now(),
        }
        
        wb = export_account_balance_report(account_data)
        response = get_excel_response('test_report', wb)
        
        assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'Content-Disposition' in response
        assert 'test_report' in response['Content-Disposition']
