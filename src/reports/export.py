"""Export utilities for generating PDF and Excel reports."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Dict, List, Optional

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def get_excel_response(filename: str, workbook: Workbook) -> HttpResponse:
    """
    Generate an HTTP response for Excel file download.
    """
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def get_pdf_response(filename: str, pdf_bytes: bytes) -> HttpResponse:
    """
    Generate an HTTP response for PDF file download.
    """
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
    return response


HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(
    start_color='1F4E78', end_color='1F4E78', fill_type='solid'
)
TITLE_FONT = Font(name='Calibri', size=14, bold=True, color='1F4E78')
TOTAL_FILL = PatternFill(
    start_color='D9D9D9', end_color='D9D9D9', fill_type='solid'
)
BORDER = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000'),
)


def set_header_row(
    worksheet,
    row_num: int,
    headers: List[str],
    col_widths: Optional[Dict[int, float]] = None,
):
    """Add formatted header row to worksheet."""
    for col_num, header_text in enumerate(headers, 1):
        cell = worksheet.cell(row=row_num, column=col_num, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=True
        )
        cell.border = BORDER

    if col_widths:
        for col_num, width in col_widths.items():
            worksheet.column_dimensions[get_column_letter(col_num)].width = (
                width
            )


def add_title_section(
    worksheet, title: str, subtitle: Optional[str] = None, start_row: int = 1
) -> int:
    """Add formatted title section to worksheet."""
    worksheet.merge_cells(f'A{start_row}:G{start_row}')
    title_cell = worksheet.cell(row=start_row, column=1, value=title)
    title_cell.font = TITLE_FONT

    current_row = start_row + 1

    if subtitle:
        worksheet.merge_cells(f'A{current_row}:G{current_row}')
        subtitle_cell = worksheet.cell(
            row=current_row, column=1, value=subtitle
        )
        subtitle_cell.font = Font(name='Calibri', size=10, bold=True)
        current_row += 1

    current_row += 1
    worksheet.cell(
        row=current_row,
        column=1,
        value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    )

    return current_row + 2


def add_total_row(
    worksheet,
    row_num: int,
    label: str,
    value: float,
    label_col: int = 1,
    value_col: int = 4,
):
    """Add formatted total row to worksheet."""
    label_cell = worksheet.cell(row=row_num, column=label_col, value=label)
    value_cell = worksheet.cell(row=row_num, column=value_col, value=value)

    for cell in [label_cell, value_cell]:
        cell.fill = TOTAL_FILL
        cell.border = BORDER
        cell.font = Font(bold=True, name='Calibri', size=11)

    value_cell.number_format = '#,##0.00'
    value_cell.alignment = Alignment(horizontal='right', vertical='center')


def export_account_balance_report(account_data: Dict[str, Any]) -> Workbook:
    """Create Excel workbook for account balance report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Account Balance"

    current_row = add_title_section(
        ws,
        "ACCOUNT BALANCE REPORT",
        f"As of {account_data['as_of_date'].strftime('%Y-%m-%d %H:%M:%S')}",
    )

    headers = ['Account Code', 'Account Name', 'Type', 'Balance']
    col_widths = {1: 15, 2: 30, 3: 15, 4: 18}
    set_header_row(ws, current_row, headers, col_widths)
    current_row += 1

    ws.cell(row=current_row, column=1, value=account_data['account_code'])
    ws.cell(row=current_row, column=2, value=account_data['account_name'])
    ws.cell(row=current_row, column=3, value=account_data['account_type'])

    balance_cell = ws.cell(
        row=current_row, column=4, value=float(account_data['balance'])
    )
    balance_cell.number_format = '#,##0.00'
    balance_cell.alignment = Alignment(horizontal='right')

    current_row += 2

    add_total_row(
        ws,
        current_row,
        'Total Debits:',
        float(account_data['debit_total']),
        value_col=2,
    )
    current_row += 1
    add_total_row(
        ws,
        current_row,
        'Total Credits:',
        float(account_data['credit_total']),
        value_col=2,
    )

    return wb


def export_transaction_history_report(
    transactions: List[Dict], account_name: str
) -> Workbook:
    """Create Excel workbook for transaction history report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction History"

    current_row = add_title_section(ws, f"TRANSACTION HISTORY - {account_name}")

    headers = [
        'Date',
        'Trans ID',
        'Description',
        'Entry Type',
        'Debit',
        'Credit',
        'Balance',
        'Status',
    ]
    col_widths = {1: 18, 2: 12, 3: 25, 4: 10, 5: 12, 6: 12, 7: 15, 8: 10}
    set_header_row(ws, current_row, headers, col_widths)
    current_row += 1

    for transaction in transactions:
        ws.cell(
            row=current_row,
            column=1,
            value=transaction['date'].strftime('%Y-%m-%d %H:%M:%S'),
        )
        ws.cell(row=current_row, column=2, value=transaction['transaction_id'])
        ws.cell(row=current_row, column=3, value=transaction['description'])
        ws.cell(row=current_row, column=4, value=transaction['entry_type'])

        debit_cell = ws.cell(
            row=current_row, column=5, value=float(transaction['debit'])
        )
        debit_cell.number_format = '#,##0.00'
        debit_cell.alignment = Alignment(horizontal='right')

        credit_cell = ws.cell(
            row=current_row, column=6, value=float(transaction['credit'])
        )
        credit_cell.number_format = '#,##0.00'
        credit_cell.alignment = Alignment(horizontal='right')

        balance_cell = ws.cell(
            row=current_row,
            column=7,
            value=float(transaction['running_balance']),
        )
        balance_cell.number_format = '#,##0.00'
        balance_cell.alignment = Alignment(horizontal='right')

        ws.cell(row=current_row, column=8, value=transaction['status'])

        current_row += 1

    return wb


def export_trial_balance_report(report_data: Dict[str, Any]) -> Workbook:
    """Create Excel workbook for trial balance report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial Balance"

    as_of_date = report_data['as_of_date'].strftime('%Y-%m-%d')
    currency_text = (
        f" - {report_data['currency']}" if report_data['currency'] else ""
    )

    current_row = add_title_section(
        ws, "TRIAL BALANCE REPORT", f"As of {as_of_date}{currency_text}"
    )

    headers = [
        'Account Code',
        'Account Name',
        'Debit Balance',
        'Credit Balance',
    ]
    col_widths = {1: 15, 2: 30, 3: 18, 4: 18}
    set_header_row(ws, current_row, headers, col_widths)
    current_row += 1

    for line in report_data['lines']:
        ws.cell(row=current_row, column=1, value=line['account_code'])
        ws.cell(row=current_row, column=2, value=line['account_name'])

        debit_cell = ws.cell(
            row=current_row, column=3, value=float(line['debit_balance'])
        )
        debit_cell.number_format = '#,##0.00'
        debit_cell.alignment = Alignment(horizontal='right')
        debit_cell.border = BORDER

        credit_cell = ws.cell(
            row=current_row, column=4, value=float(line['credit_balance'])
        )
        credit_cell.number_format = '#,##0.00'
        credit_cell.alignment = Alignment(horizontal='right')
        credit_cell.border = BORDER

        current_row += 1

    current_row += 1
    add_total_row(ws, current_row, 'TOTAL', 0, label_col=1, value_col=3)
    ws.cell(
        row=current_row, column=3, value=float(report_data['total_debits'])
    ).number_format = '#,##0.00'
    ws.cell(
        row=current_row, column=4, value=float(report_data['total_credits'])
    ).number_format = '#,##0.00'

    current_row += 2
    balance_text = (
        "✓ BALANCED" if report_data['is_balanced'] else "✗ NOT BALANCED"
    )
    balance_cell = ws.cell(row=current_row, column=1, value=balance_text)
    balance_cell.font = Font(
        bold=True, color="008000" if report_data['is_balanced'] else "FF0000"
    )

    return wb


def export_balance_sheet_report(report_data: Dict[str, Any]) -> Workbook:
    """Create Excel workbook for balance sheet report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"

    as_of_date = report_data['as_of_date'].strftime('%Y-%m-%d')
    currency_text = (
        f" - {report_data['currency']}" if report_data['currency'] else ""
    )

    current_row = add_title_section(
        ws, "BALANCE SHEET REPORT", f"As of {as_of_date}{currency_text}"
    )

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 18

    ws.cell(row=current_row, column=1, value='ASSETS').font = Font(
        size=12, bold=True, color='FFFFFF'
    )
    ws.cell(row=current_row, column=1).fill = PatternFill(
        start_color='4472C4', end_color='4472C4', fill_type='solid'
    )
    current_row += 1

    for account in report_data['assets']['accounts']:
        ws.cell(row=current_row, column=1, value=account['account_name'])
        balance_cell = ws.cell(
            row=current_row, column=2, value=float(account['balance'])
        )
        balance_cell.number_format = '#,##0.00'
        balance_cell.alignment = Alignment(horizontal='right')
        current_row += 1

    add_total_row(
        ws,
        current_row,
        'Total Assets',
        float(report_data['total_assets']),
        value_col=2,
    )
    current_row += 2

    ws.cell(row=current_row, column=1, value='LIABILITIES').font = Font(
        size=12, bold=True, color='FFFFFF'
    )
    ws.cell(row=current_row, column=1).fill = PatternFill(
        start_color='4472C4', end_color='4472C4', fill_type='solid'
    )
    current_row += 1

    total_liabilities = float(
        report_data['total_liabilities_and_equity']
    ) - float(report_data['equity']['total'])
    for account in report_data['liabilities']['accounts']:
        ws.cell(row=current_row, column=1, value=account['account_name'])
        balance_cell = ws.cell(
            row=current_row, column=2, value=float(account['balance'])
        )
        balance_cell.number_format = '#,##0.00'
        balance_cell.alignment = Alignment(horizontal='right')
        current_row += 1

    add_total_row(
        ws, current_row, 'Total Liabilities', total_liabilities, value_col=2
    )
    current_row += 2

    ws.cell(row=current_row, column=1, value='EQUITY').font = Font(
        size=12, bold=True, color='FFFFFF'
    )
    ws.cell(row=current_row, column=1).fill = PatternFill(
        start_color='4472C4', end_color='4472C4', fill_type='solid'
    )
    current_row += 1

    for account in report_data['equity']['accounts']:
        ws.cell(row=current_row, column=1, value=account['account_name'])
        balance_cell = ws.cell(
            row=current_row, column=2, value=float(account['balance'])
        )
        balance_cell.number_format = '#,##0.00'
        balance_cell.alignment = Alignment(horizontal='right')
        current_row += 1

    add_total_row(
        ws,
        current_row,
        'Total Equity',
        float(report_data['equity']['total']),
        value_col=2,
    )
    current_row += 2

    balance_text = (
        "✓ BALANCED" if report_data['is_balanced'] else "✗ NOT BALANCED"
    )
    balance_cell = ws.cell(row=current_row, column=1, value=balance_text)
    balance_color = "008000" if report_data['is_balanced'] else "FF0000"
    balance_cell.font = Font(bold=True, color=balance_color)

    return wb


def export_loan_aging_report(report_data: Dict[str, Any]) -> Workbook:
    """Create Excel workbook for loan aging report."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"

    as_of_date = report_data['as_of_date'].strftime('%Y-%m-%d')
    currency_text = (
        f" - {report_data['currency']}" if report_data['currency'] else ""
    )

    current_row = add_title_section(
        ws, "LOAN AGING REPORT - SUMMARY", f"As of {as_of_date}{currency_text}"
    )

    headers = [
        'Aging Bucket',
        'Days Range',
        'Loan Count',
        'Total Outstanding',
        'Total Principal',
        'Total Interest',
    ]
    col_widths = {1: 18, 2: 15, 3: 12, 4: 18, 5: 15, 6: 15}
    set_header_row(ws, current_row, headers, col_widths)
    current_row += 1

    for bucket in report_data['buckets']:
        ws.cell(row=current_row, column=1, value=bucket['bucket_name'])
        ws.cell(row=current_row, column=2, value=bucket['days_range'])
        ws.cell(row=current_row, column=3, value=bucket['loan_count'])

        outstanding_cell = ws.cell(
            row=current_row, column=4, value=float(bucket['total_outstanding'])
        )
        outstanding_cell.number_format = '#,##0.00'
        outstanding_cell.alignment = Alignment(horizontal='right')

        principal_cell = ws.cell(
            row=current_row, column=5, value=float(bucket['total_principal'])
        )
        principal_cell.number_format = '#,##0.00'
        principal_cell.alignment = Alignment(horizontal='right')

        interest_cell = ws.cell(
            row=current_row, column=6, value=float(bucket['total_interest'])
        )
        interest_cell.number_format = '#,##0.00'
        interest_cell.alignment = Alignment(horizontal='right')

        current_row += 1

    current_row += 1
    add_total_row(
        ws,
        current_row,
        'TOTAL',
        float(report_data['total_outstanding']),
        value_col=4,
    )

    if report_data.get('loans_detail'):
        ws_detail = wb.create_sheet("Loan Details")

        current_row = add_title_section(
            ws_detail, "LOAN AGING REPORT - DETAILS"
        )

        headers = [
            'Loan ID',
            'Borrower',
            'Principal',
            'Outstanding',
            'Interest',
            'Due Date',
            'Days Overdue',
            'Status',
        ]
        col_widths = {1: 36, 2: 20, 3: 12, 4: 15, 5: 12, 6: 12, 7: 12, 8: 10}
        set_header_row(ws_detail, current_row, headers, col_widths)
        current_row += 1

        for loan in report_data['loans_detail']:
            ws_detail.cell(
                row=current_row, column=1, value=str(loan['loan_id'])
            )
            ws_detail.cell(
                row=current_row, column=2, value=loan['borrower_name']
            )

            principal_cell = ws_detail.cell(
                row=current_row, column=3, value=float(loan['principal_amount'])
            )
            principal_cell.number_format = '#,##0.00'
            principal_cell.alignment = Alignment(horizontal='right')

            outstanding_cell = ws_detail.cell(
                row=current_row,
                column=4,
                value=float(loan['outstanding_principal']),
            )
            outstanding_cell.number_format = '#,##0.00'
            outstanding_cell.alignment = Alignment(horizontal='right')

            interest_cell = ws_detail.cell(
                row=current_row, column=5, value=float(loan['accrued_interest'])
            )
            interest_cell.number_format = '#,##0.00'
            interest_cell.alignment = Alignment(horizontal='right')

            ws_detail.cell(
                row=current_row,
                column=6,
                value=loan['due_date'].strftime('%Y-%m-%d'),
            )
            ws_detail.cell(
                row=current_row, column=7, value=loan['days_overdue']
            )
            ws_detail.cell(row=current_row, column=8, value=loan['status'])

            current_row += 1

    return wb


def export_account_balance_pdf(account_data: Dict[str, Any]) -> bytes:
    """Generate PDF for account balance report."""
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('1F4E78'),
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    story.append(Paragraph("ACCOUNT BALANCE REPORT", title_style))
    story.append(Spacer(1, 0.3 * inch))

    info = f"Account: {account_data['account_code']} - {account_data['account_name']}"
    story.append(Paragraph(info, styles['Normal']))
    story.append(
        Paragraph(
            f"As of: {account_data['as_of_date'].strftime('%Y-%m-%d %H:%M:%S')}",
            styles['Normal'],
        )
    )
    story.append(Spacer(1, 0.2 * inch))

    table_data = [
        ['Account Code', 'Account Name', 'Type', 'Balance'],
        [
            account_data['account_code'],
            account_data['account_name'],
            account_data['account_type'],
            f"${float(account_data['balance']):,.2f}",
        ],
    ]

    table = Table(
        table_data, colWidths=[1.5 * inch, 2 * inch, 1.2 * inch, 1.5 * inch]
    )
    table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('1F4E78')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 0.3 * inch))

    totals_text = f"Debit Total: ${float(account_data['debit_total']):,.2f} | Credit Total: ${float(account_data['credit_total']):,.2f}"
    story.append(Paragraph(totals_text, styles['Normal']))

    doc.build(story)
    return pdf_buffer.getvalue()


def export_transaction_history_pdf(export_data: Dict[str, Any]) -> bytes:
    """Generate PDF for transaction history report."""
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=14,
        textColor=colors.HexColor('1F4E78'),
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    story.append(Paragraph("TRANSACTION HISTORY REPORT", title_style))
    story.append(Spacer(1, 0.2 * inch))

    account_info = f"Account: {export_data['account_code']} - {export_data['account_name']}"
    story.append(Paragraph(account_info, styles['Normal']))
    story.append(Spacer(1, 0.15 * inch))

    table_data = [
        ['Transaction ID', 'Date', 'Description', 'Debit', 'Credit', 'Balance']
    ]

    for trans in export_data.get('transactions', [])[:100]:
        table_data.append(
            [
                str(trans.get('transaction_id', '')),
                (
                    trans.get('date', '').strftime('%Y-%m-%d')
                    if hasattr(trans.get('date'), 'strftime')
                    else str(trans.get('date', ''))
                ),
                str(trans.get('description', '')),
                f"${float(trans.get('debit', 0)):,.2f}",
                f"${float(trans.get('credit', 0)):,.2f}",
                f"${float(trans.get('running_balance', 0)):,.2f}",
            ]
        )

    table = Table(
        table_data,
        colWidths=[
            1.2 * inch,
            1.2 * inch,
            2.5 * inch,
            1.2 * inch,
            1.2 * inch,
            1.2 * inch,
        ],
    )
    table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('1F4E78')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('BACKGROUND', (0, 1), (-1, -1), colors.lightgrey),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('ALIGN', (3, 1), (-1, -1), 'RIGHT'),
            ]
        )
    )
    story.append(table)

    doc.build(story)
    return pdf_buffer.getvalue()


def export_trial_balance_pdf(export_data: Dict[str, Any]) -> bytes:
    """Generate PDF for trial balance report."""
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('1F4E78'),
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    story.append(Paragraph("TRIAL BALANCE REPORT", title_style))
    story.append(
        Paragraph(
            f"As of: {export_data['as_of_date'].strftime('%Y-%m-%d')}",
            styles['Normal'],
        )
    )
    story.append(Spacer(1, 0.2 * inch))

    table_data = [['Account Code', 'Account Name', 'Debit', 'Credit']]

    for account in export_data.get('accounts', []):
        table_data.append(
            [
                account['account_code'],
                account['account_name'],
                f"${float(account['debit_balance']):,.2f}",
                f"${float(account['credit_balance']):,.2f}",
            ]
        )

    table_data.append(
        [
            'TOTAL',
            '',
            f"${float(export_data['total_debits']):,.2f}",
            f"${float(export_data['total_credits']):,.2f}",
        ]
    )

    table = Table(
        table_data, colWidths=[1.5 * inch, 2.5 * inch, 1.5 * inch, 1.5 * inch]
    )
    table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('1F4E78')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 11),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('D9D9D9')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('ALIGN', (2, 1), (-1, -1), 'RIGHT'),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 0.3 * inch))

    balance_status = (
        "✓ BALANCED" if export_data['is_balanced'] else "✗ NOT BALANCED"
    )
    balance_color = colors.green if export_data['is_balanced'] else colors.red
    balance_style = ParagraphStyle(
        'BalanceStatus',
        parent=styles['Normal'],
        textColor=balance_color,
        fontSize=12,
        fontName='Helvetica-Bold',
    )
    story.append(Paragraph(balance_status, balance_style))

    doc.build(story)
    return pdf_buffer.getvalue()


def export_balance_sheet_pdf(export_data: Dict[str, Any]) -> bytes:
    """Generate PDF for balance sheet report."""
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('1F4E78'),
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    story.append(Paragraph("BALANCE SHEET REPORT", title_style))
    story.append(
        Paragraph(
            f"As of: {export_data['as_of_date'].strftime('%Y-%m-%d')}",
            styles['Normal'],
        )
    )
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("ASSETS", styles['Heading2']))
    assets_data = [['Account Code', 'Account Name', 'Balance']]
    for asset in export_data.get('assets', []):
        assets_data.append(
            [
                asset['account_code'],
                asset['account_name'],
                f"${float(asset['balance']):,.2f}",
            ]
        )
    assets_data.append(
        ['', 'Total Assets', f"${float(export_data['total_assets']):,.2f}"]
    )

    assets_table = Table(
        assets_data, colWidths=[1.5 * inch, 2.5 * inch, 1.5 * inch]
    )
    assets_table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('D9D9D9')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ]
        )
    )
    story.append(assets_table)
    story.append(Spacer(1, 0.15 * inch))

    story.append(Paragraph("LIABILITIES", styles['Heading2']))
    liab_data = [['Account Code', 'Account Name', 'Balance']]
    for liab in export_data.get('liabilities', []):
        liab_data.append(
            [
                liab['account_code'],
                liab['account_name'],
                f"${float(liab['balance']):,.2f}",
            ]
        )
    liab_data.append(
        [
            '',
            'Total Liabilities',
            f"${float(export_data['total_liabilities']):,.2f}",
        ]
    )

    liab_table = Table(
        liab_data, colWidths=[1.5 * inch, 2.5 * inch, 1.5 * inch]
    )
    liab_table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('D9D9D9')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ]
        )
    )
    story.append(liab_table)
    story.append(Spacer(1, 0.15 * inch))

    story.append(Paragraph("EQUITY", styles['Heading2']))
    eq_data = [['Account Code', 'Account Name', 'Balance']]
    for eq in export_data.get('equity', []):
        eq_data.append(
            [
                eq['account_code'],
                eq['account_name'],
                f"${float(eq['balance']):,.2f}",
            ]
        )
    eq_data.append(
        ['', 'Total Equity', f"${float(export_data['total_equity']):,.2f}"]
    )

    eq_table = Table(eq_data, colWidths=[1.5 * inch, 2.5 * inch, 1.5 * inch])
    eq_table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('D9D9D9')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ]
        )
    )
    story.append(eq_table)
    story.append(Spacer(1, 0.2 * inch))

    balance_status = (
        "✓ BALANCED" if export_data['is_balanced'] else "✗ NOT BALANCED"
    )
    balance_color = colors.green if export_data['is_balanced'] else colors.red
    balance_style = ParagraphStyle(
        'BalanceStatus',
        parent=styles['Normal'],
        textColor=balance_color,
        fontSize=12,
        fontName='Helvetica-Bold',
    )
    story.append(Paragraph(balance_status, balance_style))

    doc.build(story)
    return pdf_buffer.getvalue()


def export_loan_aging_pdf(export_data: Dict[str, Any]) -> bytes:
    """Generate PDF for loan aging report."""
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('1F4E78'),
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    story.append(Paragraph("LOAN AGING REPORT", title_style))
    story.append(
        Paragraph(
            f"As of: {export_data['as_of_date'].strftime('%Y-%m-%d')}",
            styles['Normal'],
        )
    )
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("Aging Summary", styles['Heading3']))
    summary_data = [
        [
            'Aging Bucket',
            'Days Overdue',
            'Count',
            'Principal',
            'Interest',
            'Total',
        ]
    ]

    for bucket in export_data.get('summary', []):
        summary_data.append(
            [
                bucket['bucket_name'],
                bucket['days_range'],
                str(bucket['loan_count']),
                f"${float(bucket['total_principal']):,.2f}",
                f"${float(bucket['total_interest']):,.2f}",
                f"${float(bucket['total_outstanding']):,.2f}",
            ]
        )

    summary_table = Table(
        summary_data,
        colWidths=[
            1.2 * inch,
            1.2 * inch,
            0.8 * inch,
            1.2 * inch,
            1.2 * inch,
            1.2 * inch,
        ],
    )
    summary_table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('1F4E78')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ALIGN', (3, 1), (-1, -1), 'RIGHT'),
            ]
        )
    )
    story.append(summary_table)
    story.append(Spacer(1, 0.3 * inch))

    total_outstanding = export_data.get('total_outstanding', 0)
    story.append(
        Paragraph(
            f"Total Outstanding: ${float(total_outstanding):,.2f}",
            styles['Normal'],
        )
    )

    doc.build(story)
    return pdf_buffer.getvalue()
